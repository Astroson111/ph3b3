"""Resume module — job-posting analysis, candidate profile, and the ATS
analyzer + builder.

GUARDRAIL (same discipline as the Morpheus floor): this tool ALIGNS TRUTHFUL
EXPERIENCE TO ATS VOCABULARY; IT NEVER FABRICATES QUALIFICATIONS. A missing
keyword is only inserted into a resume when it is *grounded* — i.e. the resume
already evidences that skill under different words, and the justifying line is
recorded. Keywords with no evidence are REPORTED ("you'd need to add real
experience for this"), never written into the document. It helps honest people
get read; it does not help anyone lie.

Privacy-first: a job-seeker's resume never leaves this host. All reasoning runs
on the local Hermes3 model; there are no cloud calls in the ATS path.
"""
import difflib
import json
import os
import logging
import re
import requests
from pathlib import Path

log = logging.getLogger("ph3b3.resume")

try:
    import docx  # python-docx — used only by the ATS builder's .docx output
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    log.warning("python-docx not installed — ATS resume builder unavailable (pip install python-docx)")

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False
    log.warning("beautifulsoup4 not installed — URL scraping unavailable")

try:
    import trafilatura
    TRAFILATURA_AVAILABLE = True
except ImportError:
    TRAFILATURA_AVAILABLE = False
    log.warning("trafilatura not installed — JD URL extraction falls back to BS4")

# Render-verify + relevance-weighted cutting (Ariadne v1.1). Import failures are
# survivable: the builder still produces the .docx, it just says plainly that the
# rendered page was not checked. It must never claim a clean render it did not do.
try:
    import render_verify
    RENDER_VERIFY_AVAILABLE = True
except ImportError as _e:
    RENDER_VERIFY_AVAILABLE = False
    log.warning("render_verify unavailable (%s) — built resumes will be flagged unverified", _e)

try:
    import resume_fit
    RESUME_FIT_AVAILABLE = True
except ImportError as _e:
    RESUME_FIT_AVAILABLE = False
    log.warning("resume_fit unavailable (%s) — overflow cutting disabled", _e)

# Page target for a built resume. Two pages is the ATS/recruiter convention.
DEFAULT_TARGET_PAGES = 2

# --- Ariadne JD-URL validation gate ---------------------------------------
# The fallback Ariadne returns whenever a URL can't be read/validated as a JD.
_JD_FETCH_FALLBACK = "[couldn't read that URL — paste the listing text instead.]"
# Hosts that wall content behind login / anti-bot; not worth a scraper arms race
# in v1 — short-circuit straight to the fallback.
_UNFETCHABLE_HOSTS = ("linkedin.com", "indeed.com", "glassdoor.com", "ziprecruiter.com")
# A real JD names at least a couple of these sections.
_JD_SECTION_MARKERS = (
    "responsibilit", "requirement", "qualification", "what you'll do",
    "what you will do", "who you are", "about the role", "about this role",
    "what we're looking for", "what you bring", "duties", "you will",
    "minimum qualification", "preferred qualification", "experience",
)
# Login / challenge / bot-wall fingerprints — if the extracted text reads like
# one of these, do NOT analyze it.
_BLOCK_MARKERS = (
    "enable javascript", "verify you are human", "captcha", "just a moment",
    "checking your browser", "access denied", "unusual traffic", "are you a robot",
    "sign in to continue", "log in to continue", "please enable cookies",
    "create a free account", "cloudflare",
)

_SCRAPE_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

PROFILE_PATH = Path.home() / "ph3b3_data" / "candidate_profile.json"

_EMPTY_PROFILE: dict = {
    "skills": [],
    "experience": [],
    "certifications": [],
    "projects": [],
    "notes": [],
}

# Terms whose presence in a posting is automatically worth flagging
_RED_FLAG_TERMS = [
    "rockstar", "ninja", "guru", "unicorn", "10x engineer",
    "self-starter", "fast-paced", "wear many hats", "unlimited pto",
    "competitive salary", "like a family", "startup mentality",
    "hustle", "scrappy", "unlimited growth", "culture fit",
    "work hard play hard", "entrepreneurial spirit",
]

# ---------------------------------------------------------------------------
# ATS analyzer + builder
# ---------------------------------------------------------------------------

RESUME_DIR = Path.home() / "ph3b3_data" / "resumes"   # built .docx output lives here

# Canonical ATS section -> header aliases an applicant might actually use.
# ATS parsers key off standard headers; anything not here is a "non-standard
# header" red flag (the parser may drop or misfile the section).
_SECTION_ALIASES: dict[str, list[str]] = {
    "SUMMARY":    ["summary", "professional summary", "profile", "objective",
                   "about", "about me", "career summary"],
    "EXPERIENCE": ["experience", "work experience", "professional experience",
                   "employment", "employment history", "work history", "career history"],
    "SKILLS":     ["skills", "technical skills", "core competencies", "competencies",
                   "technologies", "tech stack", "areas of expertise"],
    "EDUCATION":  ["education", "academic background", "academics"],
    # secondary sections we recognize (kept, not required for completeness)
    "PROJECTS":       ["projects", "personal projects", "selected projects", "portfolio"],
    "CERTIFICATIONS": ["certifications", "certificates", "licenses", "licenses & certifications"],
    "AWARDS":         ["awards", "honors", "achievements"],
}
# The five sections completeness is graded on:
_REQUIRED_SECTIONS = ["CONTACT", "SUMMARY", "EXPERIENCE", "SKILLS", "EDUCATION"]

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(?:\+?\d[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")


class ResumeModule:
    def __init__(self):
        self.ollama_host = os.getenv("OLLAMA_HOST", "http://localhost:11434")
        self.model = os.getenv("PH3B3_HEAVY_MODEL", os.getenv("PH3B3_MODEL", "hermes3"))
        log.info("Ariadne (resume module) ready.")
        log.info("Ariadne ATS guardrail: aligns truthful experience to ATS vocabulary; "
                 "never fabricates qualifications.")
        # Boundary (explicit): the ONLY network call Ariadne makes is a read-only
        # outbound fetch of a public JD URL the user supplies (fetch_jd). Resume
        # data is never sent anywhere — all resume reasoning is local Hermes3.

    def _is_url(self, source: str) -> bool:
        return source.strip().startswith(("http://", "https://"))

    def fetch_jd(self, url: str) -> str:
        """Ariadne's ONLY network call: a read-only, outbound fetch of the public
        JD page the user supplied. Returns extracted JD text, or _JD_FETCH_FALLBACK
        on any fetch/extraction/validation miss. Resume data is NEVER transmitted —
        this is strictly outbound to the user's URL. Best-effort on fetch-friendly
        hosts (Greenhouse/Lever/Workday/generic career pages); no headless-browser
        or anti-bot fight — login-walled sites just get the fallback."""
        url = (url or "").strip()
        if not self._is_url(url):
            return _JD_FETCH_FALLBACK
        if any(h in url.lower() for h in _UNFETCHABLE_HOSTS):
            return _JD_FETCH_FALLBACK   # login-walled / anti-bot; no arms race in v1
        try:
            log.info(f"Ariadne network call (outbound, read-only, JD fetch): {url[:90]}")
            resp = requests.get(url, headers=_SCRAPE_HEADERS, timeout=15, allow_redirects=True)
            resp.raise_for_status()
        except requests.RequestException as e:
            log.warning(f"Ariadne JD fetch failed: {e}")
            return _JD_FETCH_FALLBACK

        # Extract from raw bytes so trafilatura/BS4 detect the true encoding —
        # pages without a charset header would otherwise mojibake UTF-8 punctuation.
        raw = resp.content
        text = ""
        if TRAFILATURA_AVAILABLE:
            text = trafilatura.extract(raw, include_comments=False,
                                       include_tables=False, favor_precision=True) or ""
        if len(text) < 300 and BS4_AVAILABLE:
            text = self._bs4_extract(raw) or text            # selector-based fallback

        # validation gate: must read like a JD, must NOT read like a login/bot wall
        lower = (text or "").lower()
        if any(m in lower for m in _BLOCK_MARKERS) or not self._looks_like_jd(text):
            log.info("Ariadne JD validation gate rejected the page — returning fallback")
            return _JD_FETCH_FALLBACK
        return text.strip()

    def _bs4_extract(self, html: str) -> str:
        if not BS4_AVAILABLE:
            return ""
        soup = BeautifulSoup(html, "lxml")
        for tag in soup(["script", "style", "nav", "header", "footer", "aside", "iframe", "noscript"]):
            tag.decompose()
        for selector in [
            '[class*="job-description"]', '[class*="jobDescription"]',
            '[class*="job-details"]',    '[id*="job-description"]',
            '[class*="description"]',    "article",
            "main",                      '[role="main"]',
        ]:
            target = soup.select_one(selector)
            if target:
                text = target.get_text(separator="\n", strip=True)
                if len(text) > 200:
                    return text
        return soup.get_text(separator="\n", strip=True)

    @staticmethod
    def _looks_like_jd(text: str) -> bool:
        """Validation heuristic: long enough AND names at least two JD sections.
        400 is a floor that clears snippets/login pages while still admitting
        terse-but-real postings."""
        if not text or len(text.strip()) < 400:
            return False
        lower = text.lower()
        return sum(1 for m in _JD_SECTION_MARKERS if m in lower) >= 2

    def _auto_flags(self, text: str) -> list[str]:
        lower = text.lower()
        return [f'"{term}"' for term in _RED_FLAG_TERMS if term in lower]

    def _llm_extract(self, text: str, auto_flags: list[str]) -> str:
        excerpt = text[:4500]

        flag_note = ""
        if auto_flags:
            flag_note = (
                f"\n\nNote: the following red-flag terms were detected verbatim in the posting: "
                f"{', '.join(auto_flags)}. Make sure they appear in your RED FLAGS section."
            )

        prompt = (
            "You are a blunt, experienced technical recruiter who has read thousands of job postings "
            "and seen every corporate HR trick. Analyze the job posting below and extract each field. "
            "Do NOT invent salary information — if no salary or range is stated, write 'Not disclosed'. "
            "Treat buzzwords like 'rockstar', 'ninja', 'guru', 'passionate', 'self-starter', "
            "'fast-paced', 'wear many hats', 'unlimited PTO', 'competitive salary' (without a number), "
            "and 'like a family' as red flags, not requirements. "
            "The VERDICT must be one sentence — honest, dry, no corporate cheerleading.\n\n"
            f"JOB POSTING:\n{excerpt}"
            f"{flag_note}\n\n"
            "Respond in EXACTLY this format, nothing before or after:\n\n"
            "JOB TITLE: [title or Unknown]\n"
            "COMPANY: [company name or Unknown]\n"
            "LOCATION: [city/state, remote, hybrid, or Unknown]\n"
            "SALARY: [stated range, or 'Not disclosed']\n\n"
            "HARD REQUIREMENTS:\n"
            "- [each must-have on its own line]\n\n"
            "SOFT REQUIREMENTS:\n"
            "- [each nice-to-have on its own line, or '- None stated']\n\n"
            "RED FLAGS:\n"
            "- [specific concerns — vague language, missing salary, buzzword abuse, unrealistic scope, signs of dysfunction]\n\n"
            "CULTURE SIGNALS:\n"
            "- [what the language actually reveals between the lines]\n\n"
            "VERDICT: [one sentence]"
        )

        try:
            resp = requests.post(
                f"{self.ollama_host}/api/generate",
                json={
                    "model": self.model,
                    "prompt": prompt,
                    "stream": False,
                    "options": {"temperature": 0.3, "num_ctx": 6144},
                },
                timeout=120,
            )
            resp.raise_for_status()
            return resp.json().get("response", "").strip()
        except Exception as e:
            log.error(f"LLM extraction failed: {e}")
            return f"Extraction error: {e}"

    def extract_job_posting(self, source: str) -> str:
        source = source.strip()
        if not source:
            return "No job posting provided."

        if self._is_url(source):
            text = self.fetch_jd(source)
            if text.startswith("["):
                return text.strip("[]")   # user-facing fallback, brackets stripped
        else:
            text = source

        if len(text) < 50:
            return "Job posting text too short to analyze."

        auto_flags = self._auto_flags(text)
        return self._llm_extract(text, auto_flags)

    # ------------------------------------------------------------------
    # Profile persistence
    # ------------------------------------------------------------------

    def _load_profile(self) -> dict:
        if not PROFILE_PATH.exists():
            return {k: list(v) if isinstance(v, list) else v for k, v in _EMPTY_PROFILE.items()}
        try:
            return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as e:
            log.warning(f"Profile load failed: {e} — returning empty profile")
            return {k: list(v) if isinstance(v, list) else v for k, v in _EMPTY_PROFILE.items()}

    def _save_profile(self, profile: dict) -> None:
        PROFILE_PATH.parent.mkdir(parents=True, exist_ok=True)
        PROFILE_PATH.write_text(json.dumps(profile, indent=2, ensure_ascii=False), encoding="utf-8")

    def _format_profile(self, profile: dict) -> str:
        lines: list[str] = []

        if profile["skills"]:
            lines.append("SKILLS:\n  " + ", ".join(profile["skills"]))
        else:
            lines.append("SKILLS: (none on file)")

        lines.append("")

        if profile["experience"]:
            lines.append("EXPERIENCE:")
            for exp in profile["experience"]:
                dur = f" ({exp.get('duration', '')})" if exp.get("duration") else ""
                lines.append(f"  {exp['role']} @ {exp['org']}{dur}")
                for b in exp.get("bullets", []):
                    lines.append(f"    • {b}")
        else:
            lines.append("EXPERIENCE: (none on file)")

        lines.append("")

        if profile["certifications"]:
            lines.append("CERTIFICATIONS:")
            for cert in profile["certifications"]:
                yr = f" ({cert['year']})" if cert.get("year") else ""
                issuer = f" — {cert['issuer']}" if cert.get("issuer") else ""
                lines.append(f"  • {cert['name']}{issuer}{yr}")
        else:
            lines.append("CERTIFICATIONS: (none on file)")

        lines.append("")

        if profile["projects"]:
            lines.append("PROJECTS:")
            for proj in profile["projects"]:
                tech = f" [{', '.join(proj['tech'])}]" if proj.get("tech") else ""
                desc = f": {proj['description']}" if proj.get("description") else ""
                lines.append(f"  {proj['name']}{tech}{desc}")
                for b in proj.get("bullets", []):
                    lines.append(f"    • {b}")
        else:
            lines.append("PROJECTS: (none on file)")

        lines.append("")

        if profile["notes"]:
            lines.append("NOTES:")
            for note in profile["notes"]:
                lines.append(f"  • {note}")

        return "\n".join(lines).strip()

    # ------------------------------------------------------------------
    # Profile tool handlers
    # ------------------------------------------------------------------

    def profile_get(self) -> str:
        profile = self._load_profile()
        return self._format_profile(profile)

    def profile_add_skill(self, skill: str) -> str:
        profile = self._load_profile()
        incoming = [s.strip() for s in skill.split(",") if s.strip()]
        added = [s for s in incoming if s not in profile["skills"]]
        profile["skills"].extend(added)
        self._save_profile(profile)
        if added:
            return f"Added: {', '.join(added)}"
        return "All skills already on file — nothing added."

    def profile_add_experience(self, role: str, org: str, duration: str = "", bullets: list[str] | None = None) -> str:
        profile = self._load_profile()
        profile["experience"].append({
            "role": role,
            "org": org,
            "duration": duration,
            "bullets": bullets or [],
        })
        self._save_profile(profile)
        return f"Experience added: {role} @ {org}"

    def profile_add_certification(self, name: str, issuer: str = "", year: int | None = None) -> str:
        profile = self._load_profile()
        profile["certifications"].append({"name": name, "issuer": issuer, "year": year})
        self._save_profile(profile)
        return f"Certification added: {name}"

    def profile_add_project(self, name: str, description: str = "", tech: list[str] | None = None, bullets: list[str] | None = None) -> str:
        profile = self._load_profile()
        profile["projects"].append({
            "name": name,
            "description": description,
            "tech": tech or [],
            "bullets": bullets or [],
        })
        self._save_profile(profile)
        return f"Project added: {name}"

    def profile_add_note(self, note: str) -> str:
        profile = self._load_profile()
        profile["notes"].append(note.strip())
        self._save_profile(profile)
        return f"Note saved."

    # ------------------------------------------------------------------
    # Match engine
    # ------------------------------------------------------------------

    def match_candidate_to_job(self, job_analysis: str) -> str:
        profile = self._load_profile()
        if not any(profile[k] for k in ("skills", "experience", "certifications", "projects")):
            return (
                "Candidate profile is empty. Add skills, experience, and projects first "
                "using profile_add_skill, profile_add_experience, and profile_add_project."
            )

        profile_text = self._format_profile(profile)
        prompt = (
            "You are a technical hiring manager running an honest candidate-job match. "
            "Do not be encouraging for its own sake. Call gaps what they are.\n\n"
            f"CANDIDATE PROFILE:\n{profile_text}\n\n"
            f"JOB ANALYSIS:\n{job_analysis[:3000]}\n\n"
            "Respond in EXACTLY this format, nothing before or after:\n\n"
            "MATCH SCORE: [X]/[Y] hard requirements met\n\n"
            "GAPS:\n"
            "- [requirement]: [MET / PARTIAL / MISSING] — [one honest line]\n\n"
            "LEAD WITH:\n"
            "[One short paragraph — given this specific posting, what is the strongest angle "
            "for this candidate's application? What to emphasize at the top of the resume and in a cover note.]\n\n"
            "TAILORED BULLETS:\n"
            "- [For each hard requirement the candidate MEETS: one achievement-framed bullet "
            "grounded strictly in their actual profile. Do not invent experience.]\n\n"
            "VERDICT: [WORTH APPLYING / STRETCH / SKIP] — [one sentence reason]"
        )

        try:
            resp = requests.post(
                f"{self.ollama_host}/api/generate",
                json={
                    "model": self.model,
                    "prompt": prompt,
                    "stream": False,
                    "options": {"temperature": 0.3, "num_ctx": 8192},
                },
                timeout=150,
            )
            resp.raise_for_status()
            return resp.json().get("response", "").strip()
        except Exception as e:
            log.error(f"Match analysis failed: {e}")
            return f"Match analysis error: {e}"

    # ------------------------------------------------------------------
    # Bullet drafter
    # ------------------------------------------------------------------

    def draft_resume_section(self, requirement: str, profile_entry: str) -> str:
        prompt = (
            "Write exactly one resume bullet point that addresses the job requirement using the candidate info provided.\n\n"
            f"JOB REQUIREMENT: {requirement}\n"
            f"CANDIDATE INFO: {profile_entry}\n\n"
            "Rules — follow all of them:\n"
            "- Start with a strong past-tense action verb (Led, Built, Designed, Reduced, Migrated, Automated, etc.)\n"
            "- Include specific numbers, scale, or measurable impact if the candidate info supports it — do NOT invent them\n"
            "- No hollow adjectives (not 'successfully', not 'effectively', not 'passionate')\n"
            "- One line only — no line breaks\n"
            "- Achievement-framed: what was accomplished and what it meant, not what duties were held\n"
            "- Never start with 'Responsible for', 'Assisted with', 'Helped', or 'Worked on'\n\n"
            "Respond with ONLY the bullet text. No dash, no asterisk, no label, no preamble."
        )

        try:
            resp = requests.post(
                f"{self.ollama_host}/api/generate",
                json={
                    "model": self.model,
                    "prompt": prompt,
                    "stream": False,
                    "options": {"temperature": 0.4, "num_ctx": 4096},
                },
                timeout=60,
            )
            resp.raise_for_status()
            raw = resp.json().get("response", "").strip()
            raw = raw.lstrip("•-* ").strip()
            return f"• {raw}"
        except Exception as e:
            log.error(f"Bullet draft failed: {e}")
            return f"Draft error: {e}"

    # ==================================================================
    # ATS ANALYZER + BUILDER
    #   GUARDRAIL: aligns truthful experience to ATS vocabulary; never
    #   fabricates qualifications. Unsupported keywords are reported, not
    #   inserted. See the module docstring.
    # ==================================================================

    def _llm(self, prompt: str, temperature: float = 0.2, num_ctx: int = 8192,
             timeout: int = 150) -> str:
        """Single local-Hermes3 generate call. No cloud, ever."""
        resp = requests.post(
            f"{self.ollama_host}/api/generate",
            json={"model": self.model, "prompt": prompt, "stream": False,
                  "options": {"temperature": temperature, "num_ctx": num_ctx}},
            timeout=timeout,
        )
        resp.raise_for_status()
        return resp.json().get("response", "").strip()

    @staticmethod
    def _extract_json(raw: str, default):
        for opener, closer in (("[", "]"), ("{", "}")):
            i, j = raw.find(opener), raw.rfind(closer)
            if i != -1 and j > i:
                try:
                    return json.loads(raw[i:j + 1])
                except Exception:
                    continue
        return default

    # ---- section parsing (deterministic) -----------------------------

    @staticmethod
    def _norm_header(line: str) -> str:
        return re.sub(r"\s+", " ", line.strip().rstrip(":").strip()).lower()

    def _match_section(self, line: str) -> str | None:
        s = line.strip()
        if not s or len(s) > 40 or s.endswith((".", ",", ";")):
            return None
        norm = self._norm_header(s)
        for canon, aliases in _SECTION_ALIASES.items():
            if norm in aliases:
                return canon
        return None

    @staticmethod
    def _looks_like_header(line: str) -> bool:
        """Header-shaped: short, 1-4 words, all-caps or Title Case, not a sentence."""
        s = line.strip().rstrip(":")
        words = s.split()
        if not (1 <= len(words) <= 4) or len(s) > 40 or s.endswith((".", ",", ";")):
            return False
        if "," in s:            # a real section header never contains a comma
            return False
        letters = [c for c in s if c.isalpha()]
        if len(letters) < 3:
            return False
        return all(c.isupper() for c in letters) or s == s.title()

    def _split_sections(self, text: str):
        """Return (contact_block, {canonical: body}, order). Everything before the
        first recognized header is the contact/name block."""
        sections: dict[str, list[str]] = {}
        order: list[str] = []
        contact: list[str] = []
        current: str | None = None
        for line in text.splitlines():
            canon = self._match_section(line)
            if canon:
                current = canon
                if canon not in sections:
                    sections[canon] = []
                    order.append(canon)
                continue
            (sections[current] if current else contact).append(line)
        return ("\n".join(contact).strip(),
                {k: "\n".join(v).strip() for k, v in sections.items()},
                order)

    # ---- parse-cleanliness (deterministic) ---------------------------

    def _cleanliness(self, text: str) -> tuple[int, list[str]]:
        flags: list[str] = []
        lines = text.splitlines()

        # tables / multi-column: tabs or 2+ wide-space runs signal columns; pipes
        # only count as a table when they form a GRID (2+ pipe-heavy lines) — a
        # single "email | phone | city" contact line is normal and ATS-fine.
        pipe_lines = {i for i, ln in enumerate(lines, 1) if ln.count("|") >= 2}
        grid_pipes = pipe_lines if len(pipe_lines) >= 2 else set()
        for i, ln in enumerate(lines, 1):
            if "\t" in ln or len(re.findall(r"\S {3,}\S", ln)) >= 2 or i in grid_pipes:
                snippet = ln.strip()[:60]
                flags.append(f"line {i}: table/multi-column layout — ATS may scramble this "
                             f"(\"{snippet}\")")
                if len([f for f in flags if 'table/multi-column' in f]) >= 5:
                    flags.append("… (further table/column lines omitted)")
                    break

        # non-standard section headers
        for i, ln in enumerate(lines, 1):
            if i == 1:
                continue  # first line is the name
            if self._looks_like_header(ln) and not self._match_section(ln):
                flags.append(f"line {i}: non-standard section header \"{ln.strip()}\" — "
                             f"ATS keys off standard headers (EXPERIENCE / SKILLS / EDUCATION …)")

        # repeated content = likely header/footer bleed
        from collections import Counter
        counts = Counter(l.strip() for l in lines if 3 <= len(l.strip()) <= 60)
        for val, n in counts.items():
            if n >= 3 and not val.startswith(("•", "-", "*")):
                flags.append(f"repeated line ×{n}: \"{val}\" — looks like header/footer content "
                             f"(ATS often drops headers/footers)")

        # text-in-graphics can't be seen in plain text
        note = "text-in-graphics not detectable in plain text (re-check on .docx/.pdf source)"

        # score: start clean, penalize
        score = 100
        score -= 12 * len([f for f in flags if "table/multi-column" in f])
        score -= 8 * len([f for f in flags if "non-standard section header" in f])
        score -= 6 * len([f for f in flags if "header/footer" in f])
        score = max(0, min(100, score))
        flags.append(f"(note) {note}")
        return score, flags

    # ---- section completeness ----------------------------------------

    def _completeness(self, contact: str, sections: dict) -> tuple[list[str], list[str], list[str]]:
        present, missing, notes = [], [], []
        # CONTACT: needs an email or phone in the top block
        has_email = bool(_EMAIL_RE.search(contact))
        has_phone = bool(_PHONE_RE.search(contact))
        if has_email or has_phone:
            present.append("CONTACT")
            if not has_email:
                notes.append("CONTACT: no email detected — ATS often keys on it")
        else:
            missing.append("CONTACT")
        for sec in ("SUMMARY", "EXPERIENCE", "SKILLS", "EDUCATION"):
            if sections.get(sec):
                present.append(sec)
            else:
                missing.append(sec)
        return present, missing, notes

    # ---- keyword gap (LLM) -------------------------------------------

    def _jd_keywords(self, jd: str) -> tuple[list[str], list[str]]:
        if not jd.strip():
            return [], []
        prompt = (
            "Extract the concrete, matchable keywords from this job description — the "
            "specific tools, technologies, skills, certifications, and methodologies an "
            "applicant-tracking system would scan for. Ignore fluff and soft phrases.\n\n"
            f"JOB DESCRIPTION:\n{jd[:5000]}\n\n"
            "Respond with ONLY a JSON object, nothing else:\n"
            '{"required": ["term", ...], "preferred": ["term", ...]}\n'
            "required = must-haves; preferred = nice-to-haves. Keep each term short "
            "(a tool/skill name, not a sentence)."
        )
        try:
            data = self._extract_json(self._llm(prompt, temperature=0.1), {})
            req = [str(t).strip() for t in data.get("required", []) if str(t).strip()]
            pref = [str(t).strip() for t in data.get("preferred", []) if str(t).strip()]
            return req, pref
        except Exception as e:
            log.error(f"JD keyword extraction failed: {e}")
            return [], []

    def _classify_gaps(self, resume_text: str, missing: list[str]):
        """For terms missing verbatim from the resume, decide GROUNDED (evidenced
        under other words — with the justifying line) vs UNSUPPORTED (no evidence)."""
        if not missing:
            return [], []
        prompt = (
            "You align a resume to job keywords WITHOUT fabricating anything. For each "
            "candidate keyword below, decide if the RESUME already demonstrates that skill "
            "under different wording.\n"
            "- GROUNDED: the resume genuinely shows this skill/experience under other words. "
            "Quote the exact resume line that proves it.\n"
            "- UNSUPPORTED: there is no real evidence for it in the resume. Do NOT stretch. "
            "If unsure, mark UNSUPPORTED.\n\n"
            f"RESUME:\n{resume_text[:5000]}\n\n"
            f"KEYWORDS TO CLASSIFY: {', '.join(missing)}\n\n"
            "Respond with ONLY a JSON object, nothing else:\n"
            '{"grounded": [{"term": "...", "source": "<exact quoted resume line>"}], '
            '"unsupported": ["term", ...]}'
        )
        try:
            data = self._extract_json(self._llm(prompt, temperature=0.1), {})
            grounded, unsupported = [], []
            seen = set()
            for g in data.get("grounded", []):
                term = str(g.get("term", "")).strip()
                src = str(g.get("source", "")).strip()
                # trust-but-verify: only accept GROUNDED if the quoted source really
                # appears in the resume. Otherwise treat as unsupported.
                if term and src and self._loose_contains(resume_text, src):
                    grounded.append({"term": term, "source": src})
                    seen.add(term.lower())
                elif term:
                    unsupported.append(term)
            for t in data.get("unsupported", []):
                t = str(t).strip()
                if t and t.lower() not in seen:
                    unsupported.append(t)
            # any missing term the model dropped entirely -> unsupported (safe default)
            classified = {g["term"].lower() for g in grounded} | {u.lower() for u in unsupported}
            for m in missing:
                if m.lower() not in classified:
                    unsupported.append(m)
            return grounded, unsupported
        except Exception as e:
            log.error(f"Gap classification failed: {e}")
            return [], list(missing)

    @staticmethod
    def _loose_contains(haystack: str, needle: str) -> bool:
        norm = lambda s: re.sub(r"\s+", " ", s.lower()).strip()
        h, n = norm(haystack), norm(needle.strip('"\'')).strip()
        if len(n) < 6:
            return False
        return n in h or n[:60] in h

    @staticmethod
    def _term_in(text: str, term: str) -> bool:
        return re.search(r"\b" + re.escape(term.lower()) + r"\b", text.lower()) is not None

    # ---- public: analyzer --------------------------------------------

    def analyze_resume(self, resume_text: str, job_description: str = "",
                       source_flags: list[str] | None = None) -> str:
        if not resume_text or len(resume_text.strip()) < 40:
            return "Resume text too short to analyze — paste the full resume text."
        text = resume_text.replace("\r\n", "\n")
        contact, sections, _ = self._split_sections(text)
        score, flags = self._cleanliness(text)
        present, missing_sec, notes = self._completeness(contact, sections)

        # source_flags come from .docx/.pdf parsing (tables, text-boxes, columns,
        # header/footer content) — things invisible in flattened plain text. Each
        # knocks the cleanliness score down further.
        src = source_flags or []
        score = max(0, score - 12 * len(src))

        out = ["RESUME ANALYSIS",
               "=" * 40,
               f"\nPARSE-CLEANLINESS: {score}/100  ({'ATS-legible' if score >= 80 else 'needs work' if score >= 55 else 'high risk of misparse'})"]

        out.append("\nFORMATTING RED FLAGS:")
        real_flags = [f for f in flags if not f.startswith("(note)")]
        out += [f"  - [source] {f}" for f in src]
        out += [f"  - {f}" for f in real_flags]
        if not src and not real_flags:
            out.append("  - none detected")
        # the plain-text 'can't see graphics' caveat only applies when we had no
        # real source file to inspect
        if source_flags is None:
            out += [f"  {f}" for f in flags if f.startswith("(note)")]

        out.append("\nSECTION COMPLETENESS:")
        out.append("  present : " + (", ".join(present) if present else "none"))
        out.append("  missing : " + (", ".join(missing_sec) if missing_sec else "none — all core sections found"))
        out += [f"  ! {n}" for n in notes]

        if job_description.strip():
            req, pref = self._jd_keywords(job_description)
            all_terms = [(t, "required") for t in req] + [(t, "preferred") for t in pref]
            matched = [(t, tier) for t, tier in all_terms if self._term_in(text, t)]
            missing_terms = [t for t, _ in all_terms if not self._term_in(text, t)]
            grounded, unsupported = self._classify_gaps(text, missing_terms)

            out.append("\nKEYWORD GAP (vs job description):")
            out.append(f"  matched verbatim ({len(matched)}): " +
                       (", ".join(t for t, _ in matched) if matched else "none"))
            out.append("\n  GROUNDED — you show this under other words (safe to align):")
            out += [f"    • {g['term']}  ← your line: \"{g['source'][:90]}\"" for g in grounded] or ["    (none)"]
            out.append("\n  UNSUPPORTED — no evidence in your resume; you'd need real experience to claim these:")
            out += [f"    • {u}" for u in unsupported] or ["    (none)"]
        else:
            out.append("\nKEYWORD GAP: (paste a job description to get the keyword match)")

        return "\n".join(out)

    # ---- ATS builder (deterministic assembly, LLM only for grounding) --

    def _assemble_blocks(self, contact: str, sections: dict, grounded: list[dict]):
        """One structured representation rendered to BOTH .docx and plaintext, so
        the before/after diff always matches the produced document."""
        blocks: list[tuple[str, str]] = []
        clines = [l.strip() for l in contact.splitlines() if l.strip()]
        if clines:
            blocks.append(("name", clines[0]))
            if len(clines) > 1:
                blocks.append(("contact", " | ".join(clines[1:])))
        for canon in ("SUMMARY", "EXPERIENCE", "SKILLS", "EDUCATION",
                      "PROJECTS", "CERTIFICATIONS", "AWARDS"):
            body = sections.get(canon)
            if not body:
                continue
            if canon == "SKILLS" and grounded:
                body = self._merge_skills(body, [g["term"] for g in grounded])
            blocks.append(("header", canon))
            for ln in body.splitlines():
                t = ln.strip()
                if not t:
                    continue
                if t[0] in "•-*▪◦·":
                    blocks.append(("bullet", t.lstrip("•-*▪◦· ").strip()))
                else:
                    blocks.append(("para", t))
        return blocks

    @staticmethod
    def _merge_skills(body: str, add_terms: list[str]) -> str:
        existing = [s.strip() for s in re.split(r"[,\n]", body) if s.strip()]
        lower = {s.lower() for s in existing}
        for t in add_terms:
            if t.lower() not in lower:
                existing.append(t)
                lower.add(t.lower())
        return ", ".join(existing)

    @staticmethod
    def _blocks_to_text(blocks) -> str:
        lines = []
        for kind, val in blocks:
            if kind == "name":
                lines.append(val)
            elif kind == "contact":
                lines.append(val)
            elif kind == "header":
                lines += ["", val]
            elif kind == "bullet":
                lines.append(f"• {val}")
            else:
                lines.append(val)
        return "\n".join(lines).strip()

    def _blocks_to_docx(self, blocks, out_path: Path):
        d = docx.Document()
        for s in d.sections:            # single column, sane margins, no header/footer content
            s.top_margin = s.bottom_margin = Inches(0.6)
            s.left_margin = s.right_margin = Inches(0.7)
        normal = d.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        for kind, val in blocks:
            if kind == "name":
                p = d.add_paragraph(); r = p.add_run(val); r.bold = True; r.font.size = Pt(16)
            elif kind == "contact":
                p = d.add_paragraph(); p.add_run(val).font.size = Pt(10)
            elif kind == "header":
                p = d.add_paragraph(); p.space_before = Pt(8)
                r = p.add_run(val); r.bold = True; r.font.size = Pt(12)
            elif kind == "bullet":
                p = d.add_paragraph(f"• {val}")
                p.paragraph_format.left_indent = Inches(0.25)
            else:
                d.add_paragraph(val)
        RESUME_DIR.mkdir(parents=True, exist_ok=True)
        d.save(str(out_path))

    # Hard cap. Not a suggestion: three attempts, then hand the document over with
    # an honest warning. Never loop forever, never return nothing.
    MAX_REPAIR_ITERATIONS = 3

    @staticmethod
    def _cuttable_defects(rep) -> list[str]:
        """Which defects removing content could plausibly fix.

        Page spill, orphans and margin overflow all move when lines come out.
        FONT FALLBACK DOES NOT — no amount of cutting changes which font the
        renderer substituted. Looping on it would burn all three iterations
        achieving nothing and then warn about the same defect it started with.
        """
        return [d for d in rep.defects
                if ("spills to" in d or "orphan" in d or "bottom margin" in d)]

    def _estimate_cuts(self, blocks, rep, target_pages: int) -> int:
        """How many lines to remove this round.

        Proportional to the excess: if the render is 5 pages against a target of
        2, three fifths of the content has to go. Slight over-cut (1.15) because
        undershooting costs a whole extra render, while a marginal over-cut costs
        one bullet. Orphan-only failures need a nudge, not a haircut.
        """
        content = [i for i, (k, _) in enumerate(blocks) if k in ("bullet", "para")]
        if not content or not rep.pages:
            return 1
        if rep.pages <= target_pages:
            return 1                       # fits, but something else is off (orphan)
        excess = (rep.pages - target_pages) / rep.pages
        return max(1, int(len(content) * excess * 1.15))

    def _vision_fn(self):
        """The EXISTING vision lane (llava via local ollama), or None.

        Imported lazily and by reference so this module does not own a second
        vision path — the brief is explicit that Tier 2 reuses the take_photo /
        describe pipeline rather than building its own. Local-only: _analyze
        posts to OLLAMA_HOST on localhost, so a rasterised resume page never
        leaves this machine, which is the whole reason Ariadne exists.
        """
        try:
            import vision_module
            vm = getattr(self, "_vm", None)
            if vm is None:
                vm = self._vm = vision_module.VisionModule()
            return vm._analyze
        except Exception as e:
            log.info("[ariadne] vision lane unavailable for Tier 2: %s", e)
            return None

    def _fit_document(self, blocks, out_path: Path, target_pages: int,
                      jd_text: str, req: list[str], pref: list[str],
                      deep_verify: bool = False):
        """Render; if it does not fit, cut the lowest-scoring lines and try again.

        Returns (blocks, report, repair_log). The document is ALWAYS returned —
        a resume that overflows is still worth having; one that never arrives is
        not.
        """
        self._blocks_to_docx(blocks, out_path)
        if not RENDER_VERIFY_AVAILABLE:
            return blocks, None, []

        rep = render_verify.verify(out_path, target_pages=target_pages)
        repair_log: list[str] = []
        if not RESUME_FIT_AVAILABLE:
            return blocks, rep, repair_log

        for attempt in range(1, self.MAX_REPAIR_ITERATIONS + 1):
            fixable = self._cuttable_defects(rep)
            if rep.status != render_verify.STATUS_FAIL or not fixable:
                break                      # passed, unverifiable, or nothing cutting can fix

            n = self._estimate_cuts(blocks, rep, target_pages)
            plan = resume_fit.plan_cuts(blocks, n, req, pref, jd_text)
            if not plan.cut_indices:
                # Nothing left to cut. Stop rather than spend the remaining
                # iterations re-rendering an identical document — an iteration
                # that changes nothing is not an attempt.
                repair_log.append(f"attempt {attempt}: nothing further may be cut "
                                  f"({plan.protected_count} lines are protected) — stopping")
                break

            blocks = resume_fit.apply_cuts(blocks, plan.cut_indices)
            self._blocks_to_docx(blocks, out_path)
            before_pages = rep.pages
            rep = render_verify.verify(out_path, target_pages=target_pages)
            repair_log.append(
                f"attempt {attempt}: cut {len(plan.cut_indices)} line(s), "
                f"{before_pages} → {rep.pages} pages"
            )
            repair_log += [f"    {r}" for r in plan.rationale]
            log.info("[ariadne] repair %d/%d: cut %d, pages %s -> %s",
                     attempt, self.MAX_REPAIR_ITERATIONS, len(plan.cut_indices),
                     before_pages, rep.pages)

        # ── Tier 2 gate (verify item 6) ──────────────────────────────────────
        # The brief scopes the vision pass to "only if Tier 1 is clean, or to
        # confirm a repair", and separately requires it NOT to fire on every
        # build. Those pull against each other, because Tier 1 is clean on most
        # builds. Resolved as the intersection: Tier 1 must be clean AND there
        # must be a reason to look — either a repair just changed the document,
        # or the caller explicitly asked for a deep verify.
        #
        # So an ordinary clean build spends no vision call, a repaired document
        # gets its repair confirmed visually, and a broken document is not asked
        # about at all (Tier 1 already measured what is wrong).
        if rep is not None and rep.status == render_verify.STATUS_PASS and (repair_log or deep_verify):
            vfn = self._vision_fn()
            if vfn is not None:
                why = "confirming the repair" if repair_log else "deep verify requested"
                log.info("[ariadne] Tier 2 vision pass — %s", why)
                rep = render_verify.verify(out_path, target_pages=target_pages, vision_fn=vfn)

        return blocks, rep, repair_log

    def build_ats_resume(self, resume_text: str, job_description: str = "",
                         target_pages: int = DEFAULT_TARGET_PAGES,
                         deep_verify: bool = False) -> str:
        """Build the ATS .docx, then LOOK at the rendered page.

        target_pages is the page budget the document is held to. It is what makes
        "overflow" a defined condition rather than a feeling — before this, the
        builder emitted whatever length it emitted and nothing measured it.
        """
        if not DOCX_AVAILABLE:
            return "ATS builder unavailable — python-docx is not installed."
        try:
            target_pages = max(1, min(int(target_pages), 10))
        except (TypeError, ValueError):
            target_pages = DEFAULT_TARGET_PAGES
        if not resume_text or len(resume_text.strip()) < 40:
            return "Resume text too short to build from — paste the full resume text."
        text = resume_text.replace("\r\n", "\n")
        contact, sections, _ = self._split_sections(text)

        # grounding: only GROUNDED keywords may be inserted; unsupported are report-only
        grounded, unsupported = [], []
        if job_description.strip():
            req, pref = self._jd_keywords(job_description)
            all_terms = req + pref
            missing_terms = [t for t in all_terms if not self._term_in(text, t)]
            grounded, unsupported = self._classify_gaps(text, missing_terms)

        blocks = self._assemble_blocks(contact, sections, grounded)
        before_text = "\n".join(l.rstrip() for l in text.splitlines()).strip()

        # Fit the document to its page budget, then keep the EXACT bytes that were
        # verified. Iterating on a scratch file and copying the settled result
        # means the layout report describes the file the user downloads, rather
        # than a rebuild of it that was never rendered.
        import hashlib, shutil, tempfile
        _req = _pref = []
        if job_description.strip():
            _req, _pref = self._jd_keywords(job_description)
        _scratch = Path(tempfile.mkdtemp(prefix="ariadne-build-"))
        try:
            work = _scratch / "candidate.docx"
            blocks, rep, repair_log = self._fit_document(
                blocks, work, target_pages, job_description, _req, _pref, deep_verify)

            after_text = self._blocks_to_text(blocks)
            rid = hashlib.sha1((before_text + after_text).encode("utf-8")).hexdigest()[:12]
            out_path = RESUME_DIR / f"{rid}.docx"
            RESUME_DIR.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(work, out_path)      # byte-identical to what was verified
        finally:
            shutil.rmtree(_scratch, ignore_errors=True)

        diff = "\n".join(difflib.unified_diff(
            before_text.splitlines(), after_text.splitlines(),
            fromfile="original_resume.txt", tofile="ats_resume.docx", lineterm=""))
        if not diff.strip():
            diff = "(no textual changes — resume was already ATS-clean)"

        # ── What the rendered page actually did ──────────────────────────────
        render_line = "LAYOUT: not checked — render verification is unavailable."
        warning = ""
        if rep is not None:
            render_line = f"LAYOUT: {rep.summary()}"
            if rep.status == render_verify.STATUS_UNVERIFIED:
                log.warning("[ariadne] built %s but did not verify the render: %s", rid, rep.reason)
            elif rep.status == render_verify.STATUS_FAIL:
                # Returned anyway, per brief: never silently return a broken doc,
                # never loop forever, never return nothing. Name the defect that
                # survived so the user can act on it.
                warning = ("⚠ THIS DOCUMENT STILL HAS A LAYOUT PROBLEM after "
                           f"{len(repair_log)} repair attempt(s): "
                           + "; ".join(rep.defects)
                           + ". The file is usable and is returned regardless — "
                             "fix it by hand, raise the page target, or cut content yourself.")
                log.warning("[ariadne] %s returned with unfixed defects: %s", rid, rep.defects)

        out = ["ATS RESUME BUILT",
               "=" * 40,
               f"file: resumes/{rid}.docx   (download: GET /resume/file/{rid})",
               f"page target: {target_pages}",
               render_line]
        if warning:
            out += ["", warning]
        if rep is not None and rep.defects:
            out += [f"  ! {d}" for d in rep.defects]
        # Benign notes (e.g. the Calibri→Carlito metric-compatible substitution
        # that happens on EVERY render) are shown only when something actually
        # went wrong, where they help explain it. Printing them on a clean build
        # is two lines of noise per document telling the user nothing to act on.
        if rep is not None and rep.notes and rep.status != render_verify.STATUS_PASS:
            out += [f"  · {n}" for n in rep.notes]
        if rep is not None and rep.tier2:
            out += ["", "VISION PASS (layout only — advisory, it does not overrule the "
                        "measured checks above):"]
            out += [f"  ~ {t}" for t in rep.tier2]
        if repair_log:
            out += ["", "FITTED TO THE PAGE TARGET — what was cut and why "
                        "(overrule anything you disagree with):"]
            out += [f"  {r}" for r in repair_log]
        out += ["",
               "GROUNDED KEYWORDS INSERTED (each tied to a real line in your resume):"]
        out += [f"  • {g['term']}  ← justified by: \"{g['source'][:90]}\"" for g in grounded] \
               or ["  (none — no grounded gaps to align)"]
        out += ["",
                "REPORTED, NOT INSERTED — you'd need real experience to claim these:"]
        out += [f"  • {u}" for u in unsupported] or ["  (none)"]
        out += ["",
                "BEFORE / AFTER DIFF  (review this before using the file — it is the approval surface):",
                diff]
        return "\n".join(out)

    def get_resume_path(self, rid: str) -> Path | None:
        if not re.fullmatch(r"[0-9a-f]{12}", rid or ""):
            return None
        p = (RESUME_DIR / f"{rid}.docx").resolve()
        if p.parent == RESUME_DIR.resolve() and p.exists():
            return p
        return None

    # ---- input parsing: .txt / .docx / .pdf --------------------------

    def parse_resume_file(self, path) -> tuple[str | None, list[str]]:
        """Flatten a resume file to text + collect source-level format flags.
        Returns (text, source_flags); text is None when the file can't be parsed
        (source_flags then holds the human-readable reason)."""
        p = Path(path)
        ext = p.suffix.lower()
        if ext == ".txt":
            return p.read_text(encoding="utf-8", errors="replace"), []
        if ext == ".docx":
            if not DOCX_AVAILABLE:
                return None, ["python-docx not installed on the server"]
            try:
                return self._parse_docx(p)
            except Exception as e:
                log.error(f".docx parse failed: {e}")
                return None, [f"could not read .docx: {e}"]
        if ext == ".pdf":
            return self._parse_pdf(p)
        return None, [f"unsupported file type '{ext}' — upload .txt/.docx/.pdf or paste text"]

    def _parse_docx(self, path: Path) -> tuple[str, list[str]]:
        """Flatten in reading order (paragraphs + tables interleaved) so section
        detection works, and collect format red flags only visible at the .docx
        level (tables, columns, text-boxes/graphics, header/footer content)."""
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = Document(str(path))
        parts: list[str] = []
        for child in d.element.body.iterchildren():
            if child.tag == qn("w:p"):
                t = Paragraph(child, d).text.strip()
                if t:
                    parts.append(t)
            elif child.tag == qn("w:tbl"):
                for row in Table(child, d).rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(", ".join(dict.fromkeys(cells)))  # dedupe merged-cell repeats
        return "\n".join(parts), self._docx_flags(d)

    def _docx_flags(self, d) -> list[str]:
        from docx.oxml.ns import qn
        flags: list[str] = []
        if d.tables:
            flags.append(f"{len(d.tables)} table(s) in the source — ATS often scrambles table "
                         f"content; the rebuilt .docx flattens them to plain lines")
        for sec in d.sections:
            cols = sec._sectPr.find(qn("w:cols"))
            num = cols.get(qn("w:num")) if cols is not None else None
            if num and str(num).isdigit() and int(num) > 1:
                flags.append(f"multi-column layout ({num} columns) — ATS reads columns out of order")
                break
        xml = d.element.xml
        if "txbxContent" in xml or "<w:drawing" in xml or "<pic:pic" in xml:
            flags.append("text boxes / drawings / images detected — any text inside graphics is "
                         "invisible to most ATS")
        for sec in d.sections:
            for hf, label in ((sec.header, "header"), (sec.footer, "footer")):
                txt = " ".join(p.text for p in hf.paragraphs).strip()
                if txt:
                    flags.append(f'{label} holds content ("{txt[:40]}") — ATS frequently drops '
                                 f'headers/footers, taking that content with it')
        return flags

    def _parse_pdf(self, path: Path) -> tuple[str | None, list[str]]:
        try:
            import pdfplumber
        except ImportError:
            return None, ["pdfplumber not installed — paste text or upload .txt/.docx"]
        parts: list[str] = []
        chars = 0
        try:
            with pdfplumber.open(str(path)) as pdf:
                for pg in pdf.pages:
                    t = pg.extract_text() or ""
                    if t.strip():
                        parts.append(t)
                        chars += len(t.strip())
        except Exception as e:
            log.error(f".pdf parse failed: {e}")
            return None, [f"could not read PDF: {e}"]
        if chars < 50:
            return None, ["This looks like a scanned/image PDF — no selectable text to parse. "
                          "Export a text-based PDF, or paste your resume text instead."]
        return "\n".join(parts), []
